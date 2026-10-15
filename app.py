# -*- coding: utf-8 -*-
"""
Réconciliation de stock LEONI — application Streamlit

    Stock final = Stock initial + Mouvements      ->      Écart = SF − SI − Mouvement

Fonctionnement :
  1. choix du plant (47, 48, 66, 77)  -> table des signes du plant
  2. dépôt du stock initial et du stock final (FMALS015)
  3. dépôt de 1 ou 2 fichiers de mouvements (FMALS032)
     - le redressement CW est détecté et appliqué automatiquement, fichier par fichier
  4. choix (cases à cocher) des magasins à supprimer, parmi ceux présents dans les mouvements
  5. calcul, résultats à l'écran et export Excel

Toute la logique de calcul est dans des fonctions sans dépendance à l'interface ;
l'interface est dans main() tout en bas du fichier.
"""
from __future__ import annotations

import gc
import io
import re
from datetime import datetime, timezone

import pandas as pd
import streamlit as st

# ==============================================================================
# CONSTANTES
# ==============================================================================

LIGNES_A_IGNORER_MOUVEMENTS = 12   # lignes d'en-tête du rapport FMALS032
LIGNES_A_IGNORER_STOCK = 9         # lignes d'en-tête du rapport FMALS015

VALEURS_CW = {"CW", "CW-VW", "CW-OEM"}

# Positions (0 = colonne A) dans le fichier FMALS032 : G, I, L, M
POS_G, POS_I, POS_L, POS_M = 6, 8, 11, 12   # BookType, Warehouse, Type, by warehouses

# Signe appliqué à un type de mouvement absent de la table du plant
SIGNE_PAR_DEFAUT = 0

PLANTS = ["47", "48", "66", "77"]

# Table des signes par plant : clé = "<BookType>_<Type>", valeur = +1 / -1 / 0
SIGNES_PAR_PLANT = {
    "47": {
        "1_RECEIPT": 1,
        "12_DESPATCHED": -1,
        "13_DESPATCHED": -1,
        "16_DESPATCHED": 0,
        "17_CANCELED": 1,
        "17_DESPATCHED": -1,
        "2_RECEIPT": 0,
        "23_DESPATCHED": -1,
        "24_CANCELED": 1,
        "24_RECEIPT": 1,
        "28_UM V": 0,
        "30_RECEIPT": 1,
        "4_UM V": 0,
        "41_DESPATCHED": -1,
        "5_CANCELED": 1,
        "5_RECEIPT": 1,
        "1_CANCELED": 1,
        "16_CANCELED": 0,
        "23_": 0,
        "23_CANCELED": 1,
        "4_": 0,
        "41_CANCELED": 1,
        "13_CANCELED": 1,
        "37_DESPATCHED": -1,
        "42_UM V": 0,
        "71_RECEIPT": 0,
        "72_DESPATCHED": 0,
        "6_RECEIPT": 0,
        "2_CANCELED": 0,
    },
    "48": {
        "1_RECEIPT": 1,
        "12_DESPATCHED": -1,
        "13_DESPATCHED": -1,
        "16_DESPATCHED": 0,
        "17_CANCELED": 1,
        "17_DESPATCHED": -1,
        "2_RECEIPT": 0,
        "23_DESPATCHED": -1,
        "24_CANCELED": 1,
        "24_RECEIPT": 1,
        "28_UM V": 0,
        "30_RECEIPT": 1,
        "4_UM V": 0,
        "41_DESPATCHED": -1,
        "5_CANCELED": 1,
        "5_RECEIPT": 1,
        "1_CANCELED": 1,
        "16_CANCELED": 0,
        "23_": 0,
        "23_CANCELED": 1,
        "4_": 0,
        "41_CANCELED": 1,
        "13_CANCELED": 1,
        "37_DESPATCHED": -1,
        "42_UM V": 0,
        "36_UM V": 0,
        "20_DESPATCHED": 0,
        "2_": 0,
        "71_RECEIPT": 0,
        "72_DESPATCHED": 0,
    },
    "66": {
        "1_RECEIPT": 1,
        "12_DESPATCHED": -1,
        "13_DESPATCHED": -1,
        "16_DESPATCHED": 0,
        "17_CANCELED": 1,
        "17_DESPATCHED": -1,
        "2_RECEIPT": 0,
        "23_DESPATCHED": -1,
        "24_CANCELED": 1,
        "24_RECEIPT": 1,
        "28_UM V": 0,
        "30_RECEIPT": 1,
        "4_UM V": 0,
        "41_DESPATCHED": -1,
        "5_CANCELED": 1,
        "5_RECEIPT": 1,
        "1_CANCELED": 1,
        "16_CANCELED": 0,
        "23_": 0,
        "23_CANCELED": 1,
        "4_": 0,
        "41_CANCELED": 1,
        "13_CANCELED": 1,
        "37_DESPATCHED": 0,
        "42_UM V": 0,
        "36_UM V": -1,
        "20_DESPATCHED": 0,
        "2_": 0,
        "71_RECEIPT": 0,
        "72_DESPATCHED": 0,
        "2_CANCELED": 0,
        "28_": 0,
        "6_RECEIPT": 1,
    },
    "77": {
        "1_RECEIPT": 1,
        "12_DESPATCHED": -1,
        "13_DESPATCHED": -1,
        "16_DESPATCHED": 0,
        "17_CANCELED": 1,
        "17_DESPATCHED": -1,
        "2_RECEIPT": 0,
        "23_DESPATCHED": -1,
        "24_CANCELED": 1,
        "24_RECEIPT": 1,
        "28_UM V": 0,
        "30_RECEIPT": 1,
        "4_UM V": 0,
        "41_DESPATCHED": -1,
        "5_CANCELED": 1,
        "5_RECEIPT": 1,
        "1_CANCELED": 1,
        "16_CANCELED": 0,
        "23_": 0,
        "23_CANCELED": 1,
        "4_": 0,
        "41_CANCELED": 1,
        "13_CANCELED": 1,
        "37_DESPATCHED": -1,
        "42_UM V": 0,
        "36_UM V": -1,
        "20_DESPATCHED": 0,
        "2_": 0,
        "71_RECEIPT": 0,
        "72_DESPATCHED": 0,
        "21_RECEIPT": 0,
        "6_RECEIPT": 0,
        "2_CANCELED": 0,
    },
}


# ==============================================================================
# LECTURE ET NETTOYAGE DES FICHIERS
# ==============================================================================

def lire_csv_bytes(contenu: bytes, lignes_a_ignorer: int, nom: str = "fichier", usecols=None):
    """Lit un CSV volumineux sans créer plusieurs copies texte complètes en mémoire."""
    # Un seul nettoyage binaire du fichier. Évite decode() + split() + join() sur
    # l'intégralité d'un FMALS032 de 100 Mo ou plus.
    # Ne crée une copie complète du fichier que si des octets NUL sont réellement présents.
    contenu_propre = contenu if b"\x00" not in contenu else contenu.replace(b"\x00", b"")

    flux_entete = io.BytesIO(contenu_propre)
    entete = []
    for _ in range(lignes_a_ignorer):
        ligne = flux_entete.readline()
        if not ligne:
            raise ValueError(
                f"{nom} : moins de {lignes_a_ignorer + 1} lignes, fichier inattendu."
            )
        entete.append(ligne.decode("latin1").rstrip("\r\n"))

    try:
        df = pd.read_csv(
            io.BytesIO(contenu_propre),
            sep=";",
            skiprows=lignes_a_ignorer,
            dtype=str,
            keep_default_na=False,
            encoding="latin1",
            engine="c",
            low_memory=True,
            usecols=usecols,
        )
    except pd.errors.EmptyDataError as exc:
        raise ValueError(f"{nom} : aucune donnée après les {lignes_a_ignorer} premières lignes.") from exc

    if df.empty and len(df.columns) == 0:
        raise ValueError(f"{nom} : aucune donnée après les {lignes_a_ignorer} premières lignes.")
    return df, entete


def nettoyer_df(df: pd.DataFrame) -> pd.DataFrame:
    """Nettoie les en-têtes et uniquement les colonnes utiles au calcul/CW.

    Éviter de recopier et de nettoyer chaque cellule de fichiers de mouvements
    de plusieurs millions de lignes réduit fortement le pic de mémoire.
    """
    df.columns = [
        str(c).replace("\xa0", " ").replace("\ufeff", "").strip() for c in df.columns
    ]

    indices_utiles = {POS_G, POS_I, POS_L, POS_M}
    for nom_col in ("PartNr.", "BookType", "Quantity", "Type", "Warehouse"):
        # Nettoyer toutes les occurrences d'un nom, si un export contient des doublons.
        indices_utiles.update(i for i, c in enumerate(df.columns) if c == nom_col)
    for i in sorted(i for i in indices_utiles if 0 <= i < df.shape[1]):
        serie = df.iloc[:, i].astype(str)
        serie = serie.str.replace("\xa0", " ", regex=False).str.replace("\ufeff", "", regex=False).str.strip()
        df.iloc[:, i] = serie
    return df


def en_nombre(serie: pd.Series) -> pd.Series:
    return pd.to_numeric(
        serie.astype(str).str.replace("\xa0", "", regex=False).str.replace(",", ".", regex=False),
        errors="coerce",
    )


def normaliser_code_magasin(valeur) -> str:
    """010953 -> 10953 ; 0452CZ reste 0452CZ ; majuscules, sans espaces."""
    valeur = str(valeur).replace("\xa0", " ").strip().upper()
    if valeur in ("", "NAN", "NONE"):
        return ""
    if valeur.isdigit():
        return valeur.lstrip("0") or "0"
    return valeur


# ==============================================================================
# REDRESSEMENT CW
# ==============================================================================

def compter_cw(df: pd.DataFrame) -> int:
    """Nombre de lignes avec une valeur CW en colonne I (Warehouse) ou M (by warehouses)."""
    if df.shape[1] <= POS_M:
        return 0
    return int(
        (df.iloc[:, POS_I].isin(VALEURS_CW) | df.iloc[:, POS_M].isin(VALEURS_CW)).sum()
    )


def redresser_cw(df: pd.DataFrame):
    """
    Redressement CW :
      1. I = CW et (M vide ou CW)  -> ligne supprimée
      2. I = CW                    -> DESPATCHED : I = M, M vide, G = 17, L = DESPATCHED
      3. M = CW                    -> RECEIPT    : G = 05, L = RECEIPT, M vide
    Retourne (DataFrame redressé, statistiques).
    """
    if df.shape[1] < 13:
        raise ValueError(
            f"{df.shape[1]} colonnes seulement : au moins 13 sont nécessaires "
            "pour le redressement CW."
        )

    # Cas 1 : suppression
    col_i = df.iloc[:, POS_I]
    col_m = df.iloc[:, POS_M]
    masque_sup = col_i.isin(VALEURS_CW) & (col_m.eq("") | col_m.isin(VALEURS_CW))
    nb_supprimees = int(masque_sup.sum())
    df = df.loc[~masque_sup].copy()

    # Cas 2 : DESPATCHED
    m_i = df.iloc[:, POS_I].isin(VALEURS_CW).to_numpy()
    nb_despatched = int(m_i.sum())
    if nb_despatched:
        df.iloc[m_i, POS_I] = df.iloc[m_i, POS_M].to_numpy()
        df.iloc[m_i, POS_M] = ""
        df.iloc[m_i, POS_G] = "17"
        df.iloc[m_i, POS_L] = "DESPATCHED"

    # Cas 3 : RECEIPT
    m_m = df.iloc[:, POS_M].isin(VALEURS_CW).to_numpy()
    nb_receipt = int(m_m.sum())
    if nb_receipt:
        df.iloc[m_m, POS_G] = "05"
        df.iloc[m_m, POS_L] = "RECEIPT"
        df.iloc[m_m, POS_M] = ""

    stats = {
        "supprimees": nb_supprimees,
        "despatched": nb_despatched,
        "receipt": nb_receipt,
        "cw_restant": int(
            df.iloc[:, POS_I].isin(VALEURS_CW).sum() + df.iloc[:, POS_M].isin(VALEURS_CW).sum()
        ),
    }
    return df, stats


def exporter_csv_redresse(df: pd.DataFrame, entete: list[str]) -> bytes:
    """CSV redressé : en-tête d'origine + données redressées (mêmes conventions que le fichier source)."""
    texte = "\n".join(entete) + "\n" + df.to_csv(sep=";", index=False, lineterminator="\n")
    return texte.encode("latin1", errors="replace")


# ==============================================================================
# MOUVEMENTS (FMALS032)
# ==============================================================================

def _scanner_cw(contenu: bytes, nom: str) -> tuple[int, int, list[str]]:
    """Compte les CW en ne chargeant que I et M, sans garder le DataFrame en mémoire."""
    df, entete = lire_csv_bytes(
        contenu, LIGNES_A_IGNORER_MOUVEMENTS, nom, usecols=[POS_I, POS_M]
    )
    if df.shape[1] != 2:
        raise ValueError(f"{nom} : impossible de lire les colonnes I et M pour contrôler le CW.")
    # Les deux colonnes sélectionnées sont réordonnées en I puis M par pandas.
    for i in range(2):
        df.iloc[:, i] = (
            df.iloc[:, i].astype(str).str.replace("\xa0", " ", regex=False).str.strip()
        )
    nb_cw = int((df.iloc[:, 0].isin(VALEURS_CW) | df.iloc[:, 1].isin(VALEURS_CW)).sum())
    nb_lues = len(df)
    del df
    return nb_lues, nb_cw, entete


def _parse_mouvements(contenu: bytes, nom: str) -> dict:
    """Analyse en deux passes : charge seulement 5 colonnes si aucun CW n'est présent."""
    nb_lues, nb_cw, entete = _scanner_cw(contenu, nom)
    if nb_lues == 0:
        raise ValueError(f"{nom} : aucune ligne de mouvement.")

    redresse, stats, csv_redresse = False, None, None
    if nb_cw > 0:
        # Le redressement CW dépend des positions G, I, L et M : charger toutes les colonnes
        # uniquement si nécessaire, puis sauvegarder l'export avant de réduire le DataFrame.
        df, _ = lire_csv_bytes(contenu, LIGNES_A_IGNORER_MOUVEMENTS, nom)
        df = nettoyer_df(df)
        manquantes = [c for c in ("PartNr.", "BookType", "Quantity", "Type") if c not in df.columns]
        if manquantes:
            raise ValueError(
                f"{nom} : colonnes manquantes {manquantes}. Est-ce bien un fichier FMALS032 ?"
            )
        df, stats = redresser_cw(df)
        csv_redresse = exporter_csv_redresse(df, entete)
        redresse = True
        # Les 25+ colonnes ne sont plus nécessaires aux calculs après le redressement.
        # Garder uniquement les 5 colonnes utilisées par les magasins et le TCD.
        noms_utiles = ("PartNr.", "BookType", "Warehouse", "Quantity", "Type")
        manquantes_utiles = [c for c in noms_utiles if c not in df.columns]
        if manquantes_utiles:
            raise ValueError(f"{nom} : colonnes utiles absentes après CW : {manquantes_utiles}")
        df = df.loc[:, list(noms_utiles)].copy()
        gc.collect()
    else:
        # Sans CW, les colonnes descriptives ne servent ni à la réconciliation ni au TCD.
        # Sélection par position pour tolérer les espaces dans les en-têtes exportés.
        df, _ = lire_csv_bytes(
            contenu, LIGNES_A_IGNORER_MOUVEMENTS, nom,
            usecols=[2, 6, 8, 10, 11],
        )
        if df.shape[1] != 5:
            raise ValueError(f"{nom} : structure inattendue, 5 colonnes utiles attendues.")
        noms = [str(c).replace("\xa0", " ").replace("\ufeff", "").strip() for c in df.columns]
        df.columns = noms
        # Le rapport peut contenir des espaces en fin de nom (ex. 'BookType ').
        mapping = {}
        for idx, canonical in zip(range(5), ("PartNr.", "BookType", "Warehouse", "Quantity", "Type")):
            mapping[df.columns[idx]] = canonical
        df.rename(columns=mapping, inplace=True)
        manquantes = [c for c in ("PartNr.", "BookType", "Quantity", "Type", "Warehouse") if c not in df.columns]
        if manquantes:
            raise ValueError(
                f"{nom} : colonnes utiles introuvables {manquantes}. Vérifier la structure FMALS032."
            )
        for col in ("PartNr.", "BookType", "Warehouse", "Quantity", "Type"):
            df[col] = df[col].astype(str).str.replace("\xa0", " ", regex=False).str.strip()

    return {
        "nom": nom,
        "df": df,
        "entete": entete,
        "nb_lues": nb_lues,
        "nb_cw": nb_cw,
        "redresse": redresse,
        "stats": stats,
        "csv_redresse": csv_redresse,
    }


def colonne_warehouse(df: pd.DataFrame) -> pd.Series:
    """Colonne Warehouse (par son nom si unique, sinon colonne I)."""
    if list(df.columns).count("Warehouse") == 1:
        return df["Warehouse"]
    if df.shape[1] > POS_I:
        return df.iloc[:, POS_I]
    raise ValueError("Colonne 'Warehouse' introuvable dans le fichier de mouvements.")


def lister_magasins(dfs: list[pd.DataFrame]) -> pd.DataFrame:
    """
    Magasins présents dans les fichiers de mouvements (après redressement CW).
    Index = code normalisé ; colonnes : Magasin (libellé), Lignes (nombre de lignes).
    """
    series = [colonne_warehouse(df) for df in dfs]
    s = pd.concat(series, ignore_index=True)
    s = s[s != ""]
    if s.empty:
        return pd.DataFrame({"Magasin": [], "Lignes": []}, index=pd.Index([], name="code"))

    tab = pd.DataFrame({"code": s.map(normaliser_code_magasin).to_numpy(), "Magasin": s.to_numpy()})
    out = tab.groupby("code").agg(
        Magasin=("Magasin", lambda x: " / ".join(sorted(set(x)))),
        Lignes=("Magasin", "size"),
    )
    return out.sort_index()


def supprimer_magasins(df: pd.DataFrame, codes: set[str]):
    """Supprime les lignes dont le magasin (normalisé) est dans `codes`. Retourne (df, nb_supprimées)."""
    if not codes:
        return df, 0
    masque = colonne_warehouse(df).map(normaliser_code_magasin).isin(codes)
    return df.loc[~masque], int(masque.sum())


def creer_tcd(df: pd.DataFrame, signes: dict, nom: str):
    """
    TCD : PartNumber en lignes, type de mouvement (BookType_Type) en colonnes,
    signes du plant appliqués + colonne Total.
    Retourne (tcd, tableau des mouvements inconnus).
    """
    requises = ["PartNr.", "BookType", "Quantity", "Type"]
    manquantes = [c for c in requises if c not in df.columns]
    if manquantes:
        raise ValueError(f"{nom} : colonnes manquantes {manquantes}")

    # Une colonne requise en double : on ne devine pas laquelle utiliser
    doublons = [c for c in requises if (df.columns == c).sum() > 1]
    if doublons:
        raise ValueError(
            f"{nom} : colonne(s) requise(s) en double {doublons}. "
            "Vérifier les en-têtes du fichier (ex. 'Type' et 'Type ')."
        )
    df = df.loc[:, ~df.columns.duplicated()]

    d = df[requises].copy()
    d["Quantity"] = en_nombre(d["Quantity"]).fillna(0)
    d["BookType_normalise"] = d["BookType"].str.replace(r"^0+(\d+)$", r"\1", regex=True)
    d["Mouv"] = d["BookType_normalise"] + "_" + d["Type"]

    # Mouvements absents de la table des signes
    inconnus_noms = sorted(set(d["Mouv"]) - set(signes))
    if inconnus_noms:
        inconnus = (
            d[d["Mouv"].isin(inconnus_noms)]
            .groupby("Mouv")["Quantity"].agg(["count", "sum"])
            .reset_index()
            .rename(columns={"Mouv": "Mouvement", "count": "Lignes", "sum": "Quantité"})
        )
        inconnus.insert(0, "Fichier", nom)
    else:
        inconnus = pd.DataFrame(columns=["Fichier", "Mouvement", "Lignes", "Quantité"])

    if d.empty:
        return pd.DataFrame({"PartNumber": [], "Total": []}), inconnus

    tcd = pd.pivot_table(
        d, index="PartNr.", columns="Mouv", values="Quantity", aggfunc="sum", fill_value=0,
    ).reset_index().rename(columns={"PartNr.": "PartNumber"})

    colonnes = [c for c in tcd.columns if c != "PartNumber"]
    for col in colonnes:
        # Signe absent de la table -> SIGNE_PAR_DEFAUT (0)
        tcd[col] = tcd[col] * signes.get(str(col).strip(), SIGNE_PAR_DEFAUT)
    tcd["Total"] = tcd[colonnes].sum(axis=1)
    return tcd, inconnus


# ==============================================================================
# STOCKS (FMALS015)
# ==============================================================================

def _parse_stock(contenu: bytes, nom: str) -> pd.DataFrame:
    # Pour FMALS015, le calcul utilise uniquement les colonnes A, D, G et J.
    # Lire seulement A:J réduit fortement la mémoire pour les fichiers de stock.
    df, _ = lire_csv_bytes(contenu, LIGNES_A_IGNORER_STOCK, nom, usecols=range(10))
    df = nettoyer_df(df)
    if df.shape[1] < 10:
        raise ValueError(
            f"{nom} : {df.shape[1]} colonnes seulement, au moins 10 sont attendues "
            "(fichier FMALS015 ?)."
        )

    stock = en_nombre(df.iloc[:, 6])
    value = en_nombre(df.iloc[:, 9])
    df["PU"] = (value / stock).replace([float("inf"), -float("inf")], 0).fillna(0)

    df["_PartNumber"] = df.iloc[:, 0].astype(str).str.strip()
    df["_Stock"] = stock.fillna(0)
    return df[df["_PartNumber"] != ""].copy()


# Les fichiers sont lus une seule fois par contenu : changer une case à cocher ne relit rien.
# Le cache est limité aux derniers fichiers et expire après 15 minutes.
@st.cache_resource(show_spinner=False, ttl=900, max_entries=2)
def parse_mouvements(contenu: bytes, nom: str) -> dict:
    return _parse_mouvements(contenu, nom)


@st.cache_resource(show_spinner=False, ttl=900, max_entries=2)
def parse_stock(contenu: bytes, nom: str) -> pd.DataFrame:
    return _parse_stock(contenu, nom)


# ==============================================================================
# RÉCONCILIATION
# ==============================================================================

def _premier_par_part(df: pd.DataFrame, valeurs: pd.Series) -> pd.Series:
    tmp = df.assign(_v=valeurs).drop_duplicates("_PartNumber")
    return tmp.set_index("_PartNumber")["_v"]


def reconcilier(df_si: pd.DataFrame, df_sf: pd.DataFrame, tcds: list[pd.DataFrame]) -> pd.DataFrame:
    """Une ligne par PartNumber (union SI + SF) ; Écart = SF − SI − Mouvement."""
    deux = len(tcds) == 2
    parts = sorted(set(df_si["_PartNumber"]) | set(df_sf["_PartNumber"]))
    rec = pd.DataFrame({"PartNumber": parts})

    rec["SI"] = rec["PartNumber"].map(df_si.groupby("_PartNumber")["_Stock"].sum()).fillna(0)
    rec["SF"] = rec["PartNumber"].map(df_sf.groupby("_PartNumber")["_Stock"].sum()).fillna(0)

    colonnes_mouv = []
    for i, tcd in enumerate(tcds, start=1):
        nom_col = f"Mouvement {i}" if deux else "Mouvement"
        rec[nom_col] = rec["PartNumber"].map(tcd.set_index("PartNumber")["Total"]).fillna(0)
        colonnes_mouv.append(nom_col)
    if deux:
        rec["Mouvement"] = rec["Mouvement 1"] + rec["Mouvement 2"]

    # Type (colonne D) et PU : valeur du SF en priorité, sinon du SI
    type_sf = _premier_par_part(df_sf, df_sf.iloc[:, 3])
    type_si = _premier_par_part(df_si, df_si.iloc[:, 3])
    pu_sf = _premier_par_part(df_sf, df_sf["PU"])
    pu_si = _premier_par_part(df_si, df_si["PU"])

    rec["Type"] = rec["PartNumber"].map(type_sf).fillna(rec["PartNumber"].map(type_si)).fillna("")
    rec["PU"] = pd.to_numeric(
        rec["PartNumber"].map(pu_sf).fillna(rec["PartNumber"].map(pu_si)), errors="coerce"
    ).fillna(0)

    # Arrondi à 6 décimales : supprime le bruit de calcul flottant (ex. 1e-13)
    rec["Ecart"] = (rec["SF"] - rec["SI"] - rec["Mouvement"]).round(6)
    rec["Ecart €"] = rec["Ecart"] * rec["PU"]

    ordre = ["PartNumber", "Type", "PU", "SI"] + colonnes_mouv
    if deux:
        ordre.append("Mouvement")
    ordre += ["SF", "Ecart", "Ecart €"]
    return rec[ordre]


def trouver_orphelins(df_rec: pd.DataFrame, tcds: list[pd.DataFrame], noms: list[str]) -> pd.DataFrame:
    """
    PartNumbers qui ont des mouvements mais n'existent ni dans le stock initial ni dans
    le stock final : ils ne figurent pas dans la réconciliation.
    """
    connus = set(df_rec["PartNumber"])
    lignes = []
    for tcd, nom in zip(tcds, noms):
        o = tcd[~tcd["PartNumber"].isin(connus) & (tcd["Total"].abs() > 1e-9)]
        for p, t in zip(o["PartNumber"], o["Total"]):
            lignes.append({"Fichier": nom, "PartNumber": p, "Mouvement": t})
    return pd.DataFrame(lignes, columns=["Fichier", "PartNumber", "Mouvement"])


def calculer(plant: str, df_si, nom_si: str, df_sf, nom_sf: str, mouvs: list[dict], codes: set[str],
             libelles_magasins: dict | None = None) -> dict:
    """Calcul complet. Retourne un dictionnaire avec tous les résultats et le fichier Excel."""
    signes = SIGNES_PAR_PLANT[plant]
    tcds, inconnus_liste, controle = [], [], []
    csv_redresses = []

    for m in mouvs:
        df, nb_sup_mag = supprimer_magasins(m["df"], codes)
        tcd, inconnus = creer_tcd(df, signes, m["nom"])
        tcds.append(tcd)
        inconnus_liste.append(inconnus)

        stats = m["stats"] or {}
        controle.append({
            "Fichier": m["nom"],
            "Lignes lues": m["nb_lues"],
            "Lignes CW détectées": m["nb_cw"],
            "Redressement CW": "appliqué" if m["redresse"] else "non nécessaire",
            "Lignes supprimées (CW)": stats.get("supprimees", 0),
            "Lignes supprimées (magasins)": nb_sup_mag,
            "Lignes retenues": len(df),
        })
        if m["redresse"]:
            csv_redresses.append((m["nom"], m["csv_redresse"]))

    inconnus_all = pd.concat(inconnus_liste, ignore_index=True) if inconnus_liste else pd.DataFrame()
    rec = reconcilier(df_si, df_sf, tcds)
    orphelins = trouver_orphelins(rec, tcds, [m["nom"] for m in mouvs])

    libelles = libelles_magasins or {}
    magasins_supprimes = [libelles.get(c, c) for c in sorted(codes)]

    params = {
        "Plant": plant,
        "Généré le (UTC)": datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M"),
        "Stock initial": nom_si,
        "Stock final": nom_sf,
        "Fichiers de mouvements": " | ".join(m["nom"] for m in mouvs),
        "Redressement CW": " | ".join(
            f"{m['nom']} : {'appliqué (' + str(m['nb_cw']) + ' lignes CW)' if m['redresse'] else 'non nécessaire'}"
            for m in mouvs
        ),
        "Magasins supprimés": ", ".join(magasins_supprimes) if magasins_supprimes else "aucun",
        "Mouvements inconnus (signe 0)": int(len(inconnus_all)),
        "Formule": "Ecart = SF - SI - Mouvement ; Ecart EUR = Ecart x PU",
    }
    controle_df = pd.DataFrame(controle)

    # Construire le fichier Excel une seule fois ; le résultat conserve uniquement
    # les objets nécessaires à l'affichage et aux boutons de téléchargement.
    excel_bytes = construire_excel(rec, tcds, params, inconnus_all, orphelins)
    resultat = {
        "plant": plant,
        "df": rec,
        "tcds": tcds,
        "noms_mouvements": [m["nom"] for m in mouvs],
        "inconnus": inconnus_all,
        "controle": controle_df,
        "orphelins": orphelins,
        "magasins_supprimes": magasins_supprimes,
        "csv_redresses": csv_redresses,
        "params": params,
        "excel": excel_bytes,
        "nom_excel": f"Reconciliation_Plant{plant}_{datetime.now(timezone.utc).strftime('%Y%m%d_%H%M')}.xlsx",
    }

    # Libérer les objets intermédiaires locaux dès qu'ils ne sont plus nécessaires.
    del controle, controle_df, excel_bytes, inconnus_liste, libelles, params, rec, orphelins
    gc.collect()
    return resultat


# ==============================================================================
# EXPORT EXCEL
# ==============================================================================

def _styler_feuille(ws, df: pd.DataFrame, colonnes_nombres: tuple = (), colonnes_ecart: tuple = ()):
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    fond = PatternFill("solid", fgColor="1F3A5F")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fond
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    if len(df.columns) and len(df):
        ws.auto_filter.ref = ws.dimensions

    n = len(df)
    for idx, col in enumerate(df.columns, start=1):
        lettre = get_column_letter(idx)
        largeur = max([len(str(col))] + [len(str(v)) for v in df[col].head(500)])
        ws.column_dimensions[lettre].width = min(max(largeur + 2, 10), 42)
        if col in colonnes_nombres and n:
            for cell in ws[lettre][1:]:
                cell.number_format = "#,##0.00"
        if col in colonnes_ecart and n:
            ws.conditional_formatting.add(
                f"{lettre}2:{lettre}{n + 1}",
                CellIsRule(operator="notEqual", formula=["0"], font=Font(color="C00000", bold=True)),
            )


def construire_excel(rec, tcds, params, inconnus, orphelins) -> bytes:
    """Classeur : Réconcil, TCD1 (TCD2), Paramètres, + Mouvements inconnus / Hors stock si besoin."""
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        rec.to_excel(writer, sheet_name="Réconcil", index=False)
        _styler_feuille(
            writer.sheets["Réconcil"], rec,
            colonnes_nombres=("PU", "Ecart €"), colonnes_ecart=("Ecart", "Ecart €"),
        )
        for i, tcd in enumerate(tcds, start=1):
            tcd.to_excel(writer, sheet_name=f"TCD{i}", index=False)
            _styler_feuille(writer.sheets[f"TCD{i}"], tcd)

        p = pd.DataFrame({"Paramètre": list(params), "Valeur": [str(v) for v in params.values()]})
        p.to_excel(writer, sheet_name="Paramètres", index=False)
        _styler_feuille(writer.sheets["Paramètres"], p)
        writer.sheets["Paramètres"].column_dimensions["B"].width = 90

        if len(inconnus):
            inconnus.to_excel(writer, sheet_name="Mouvements inconnus", index=False)
            _styler_feuille(writer.sheets["Mouvements inconnus"], inconnus)
        if len(orphelins):
            orphelins.to_excel(writer, sheet_name="Hors stock", index=False)
            _styler_feuille(writer.sheets["Hors stock"], orphelins)
    return buf.getvalue()


# ==============================================================================
# INTERFACE
# ==============================================================================

CSS = """
<style>
:root{--navy:#002857;--navy2:#073d78;--orange:#ff7514;--ink:#17283f;--muted:#6d7c91;--line:#dce5ef;--paper:#f4f7fb;--white:#fff}
.stApp{background:radial-gradient(ellipse at 88% 0%,rgba(255,117,20,.055),transparent 28%),var(--paper);color:var(--ink)}
.block-container{max-width:1460px;padding-top:1.1rem;padding-bottom:4rem}
[data-testid="stHeader"]{background:rgba(244,247,251,.86)}
#MainMenu,footer{visibility:hidden}
/* Brand-led editorial hero */
.hero{position:relative;isolation:isolate;overflow:hidden;display:flex;align-items:center;justify-content:space-between;gap:2rem;min-height:228px;padding:2.1rem 2.4rem;margin:.25rem 0 1.6rem;border-radius:24px;background:linear-gradient(115deg,#001b3b 0%,#002857 52%,#084985 100%);box-shadow:0 18px 44px rgba(0,40,87,.16);color:white}
.hero:before{content:"";position:absolute;z-index:-1;width:330px;height:330px;right:8%;top:-190px;border:1px solid rgba(255,255,255,.17);border-radius:50%;box-shadow:0 0 0 32px rgba(255,255,255,.025),0 0 0 66px rgba(255,255,255,.025),0 0 0 100px rgba(255,255,255,.02)}
.hero:after{content:"";position:absolute;z-index:-1;left:0;bottom:0;width:38%;height:5px;background:linear-gradient(90deg,var(--orange),#ffb16f,transparent)}
.hero-main{max-width:760px}.hero .eyebrow{display:flex;align-items:center;gap:.55rem;color:#ffc18f;font-size:.72rem;font-weight:800;letter-spacing:.17em;text-transform:uppercase;margin:0 0 1rem}.hero .eyebrow:before{content:"";width:23px;height:3px;background:var(--orange);border-radius:4px}
.hero h1{color:#fff;margin:0;padding:0;font-size:clamp(2rem,3.6vw,3.15rem);font-weight:800;letter-spacing:-.055em;line-height:1.04}.hero p{max-width:680px;color:#d9e5f2;font-size:1rem;line-height:1.65;margin:.9rem 0 0}.hero .formula{display:inline-flex;align-items:center;gap:.5rem;margin-top:1.2rem;padding:.6rem .85rem;border:1px solid rgba(255,255,255,.18);border-radius:9px;background:rgba(255,255,255,.07);font:600 .79rem ui-monospace,SFMono-Regular,Menlo,monospace;color:#fff}.hero .hero-side{flex:0 0 180px;align-self:stretch;display:flex;flex-direction:column;justify-content:center;align-items:flex-end;text-align:right}.hero .side-kicker{font-size:.65rem;letter-spacing:.19em;text-transform:uppercase;color:#9db6d2;font-weight:800}.hero .side-number{font-size:5.3rem;line-height:.95;font-weight:850;letter-spacing:-.09em;color:rgba(255,255,255,.94)}.hero .side-number span{color:var(--orange)}.hero .side-label{margin-top:.65rem;font-size:.78rem;color:#c6d6e8;max-width:130px}
/* Sidebar as control room */
section[data-testid="stSidebar"]{background:#fff;border-right:1px solid var(--line)}
section[data-testid="stSidebar"]>div{padding-top:1.35rem}
section[data-testid="stSidebar"] h3{color:var(--navy);font-weight:800;letter-spacing:-.025em}
section[data-testid="stSidebar"] [data-testid="stMarkdownContainer"] p{color:#52647a;line-height:1.6;font-size:.9rem}
section[data-testid="stSidebar"] hr{border-color:#e5ebf3}
.workflow-kicker{margin:.3rem 0 .8rem;color:#7a8aa0;font-size:.69rem;font-weight:850;letter-spacing:.18em}
/* Main workflow cards */
div[data-testid="stVerticalBlockBorderWrapper"]{background:rgba(255,255,255,.94);border:1px solid #e0e7f0;border-radius:17px;box-shadow:0 6px 22px rgba(0,40,87,.035);padding:.35rem .35rem;transition:border-color .2s,box-shadow .2s}
div[data-testid="stVerticalBlockBorderWrapper"]:hover{border-color:#c8d6e6;box-shadow:0 9px 27px rgba(0,40,87,.055)}
div[data-testid="stVerticalBlockBorderWrapper"] h4{color:var(--navy);font-weight:800;font-size:1.06rem;letter-spacing:-.025em;padding-bottom:.55rem;border-bottom:1px solid #edf1f6}
/* Inputs and drag-drop */
.stSelectbox label,.stFileUploader label,.stTextInput label,.stToggle label{color:var(--navy);font-weight:700;font-size:.9rem}
div[data-testid="stFileUploader"] section{background:linear-gradient(145deg,#fbfcfe,#f5f8fc);border:1.5px dashed #b9c9dc;border-radius:13px;transition:all .18s}
div[data-testid="stFileUploader"] section:hover{border-color:var(--orange);background:#fff8f2}
.stSelectbox [data-baseweb="select"]>div,.stTextInput input{border-radius:10px;border-color:#d6e0ec}
/* Buttons with deliberate hierarchy */
.stButton button,.stDownloadButton button{border-radius:10px!important;font-weight:750!important;min-height:2.7rem;transition:transform .16s,box-shadow .16s,background .16s}
.stButton button[kind="primary"],.stDownloadButton button[kind="primary"]{background:var(--orange)!important;border:1px solid var(--orange)!important;color:#fff!important;box-shadow:0 5px 12px rgba(255,117,20,.17)}
.stButton button[kind="primary"]:hover,.stDownloadButton button[kind="primary"]:hover{background:#e96205!important;border-color:#e96205!important;transform:translateY(-1px);box-shadow:0 8px 18px rgba(255,117,20,.24)}
.stButton button:not([kind="primary"]),.stDownloadButton button:not([kind="primary"]){border:1px solid #cfdae7;color:var(--navy);background:#fff}
/* Result metrics: dashboard tiles */
div[data-testid="stMetric"]{position:relative;overflow:hidden;background:#fff;border:1px solid #e1e8f1;border-radius:15px;padding:1rem 1rem .9rem;box-shadow:0 5px 18px rgba(0,40,87,.035)}
div[data-testid="stMetric"]:before{content:"";position:absolute;left:0;top:0;bottom:0;width:4px;background:var(--orange)}
div[data-testid="stMetricLabel"]{color:var(--muted);font-weight:700;font-size:.78rem}
div[data-testid="stMetricValue"]{color:var(--navy);font-weight:850;letter-spacing:-.04em}
/* Tabs, tables, alerts */
button[data-baseweb="tab"]{color:#65758a;font-weight:750;padding-top:.8rem;padding-bottom:.8rem}
button[data-baseweb="tab"][aria-selected="true"]{color:var(--navy)}
div[data-baseweb="tab-highlight"]{background:var(--orange);height:3px}
div[data-testid="stDataFrame"],div[data-testid="stDataEditor"]{border:1px solid #dfe7f0;border-radius:12px;overflow:hidden}
.stAlert{border-radius:12px;border:1px solid #e4eaf2}
hr{border-color:#e2e9f1}
/* Primary run action */
.stButton:has(button[kind="primary"]){padding-top:.35rem}
@media(max-width:760px){.block-container{padding-top:.7rem}.hero{padding:1.5rem 1.3rem;min-height:0;border-radius:18px}.hero .hero-side{display:none}.hero h1{font-size:2rem}.hero p{font-size:.92rem}.hero .formula{font-size:.69rem;line-height:1.5;white-space:normal}div[data-testid="stVerticalBlockBorderWrapper"]{padding:.15rem}}
</style>
"""

HERO = (
    '<div class="hero">'
    '<div class="hero-main">'
    '<div class="eyebrow">LEONI · Supply Chain Management</div>'
    '<h1>Stock, sous contrôle.</h1>'
    '<p>Une vision fiable des mouvements, des stocks et des écarts. Importez vos fichiers, contrôlez les paramètres, puis lancez une réconciliation traçable.</p>'
    '<span class="formula">ÉCART = STOCK FINAL − STOCK INITIAL − MOUVEMENTS</span>'
    '</div>'
    '<div class="hero-side"><div class="side-kicker">Stock intelligence</div>'
    '<div class="side-number">S<span>.</span></div>'
    '<div class="side-label">Réconciliation & contrôle</div></div>'
    '</div>'
)


def _fmt(x: float, decimales: int = 2) -> str:
    return f"{x:,.{decimales}f}".replace(",", " ")


def _fid(f) -> tuple:
    """Identité d'un fichier déposé (nom, taille, identifiant de dépôt)."""
    return (f.name, f.size, getattr(f, "file_id", None))


def plant_dans_nom(nom: str):
    """'FMALS015_W66_20260901.csv' -> '66' (None si absent)."""
    m = re.search(r"W(\d{2})", nom.upper())
    return m.group(1) if m else None


def table_signes(plant: str) -> pd.DataFrame:
    signes = SIGNES_PAR_PLANT[plant]
    effet = {1: "Entrée (+)", -1: "Sortie (−)", 0: "Neutre (0)"}
    return pd.DataFrame({
        "Mouvement": list(signes),
        "Signe": list(signes.values()),
        "Effet": [effet[v] for v in signes.values()],
    })


def afficher_resultats(res: dict):
    df = res["df"]
    nb_ecart = int((df["Ecart"].abs() > 1e-9).sum())

    st.markdown(f"### Résultats — plant {res['plant']}")
    k1, k2, k3, k4, k5 = st.columns(5)
    k1.metric("PartNumbers", _fmt(len(df), 0))
    k2.metric("Articles avec écart", _fmt(nb_ecart, 0))
    k3.metric("Écart net (quantité)", _fmt(df["Ecart"].sum()))
    k4.metric("Écart net (€)", _fmt(df["Ecart €"].sum()))
    k5.metric("Écart absolu (€)", _fmt(df["Ecart €"].abs().sum()))

    st.download_button(
        "⬇️ Télécharger le fichier Excel",
        data=res["excel"],
        file_name=res["nom_excel"],
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        type="primary",
    )

    t_rec, t_top, t_tcd, t_ctrl = st.tabs(
        ["📋 Réconciliation", "📊 Plus gros écarts", "🔀 TCD", "🔍 Contrôles"]
    )

    with t_rec:
        c1, c2 = st.columns([3, 1])
        recherche = c1.text_input("Rechercher un PartNumber", placeholder="ex. 12345", key="recherche")
        seulement = c2.toggle("Seulement les écarts", value=False, key="seulement_ecarts")
        vue = df
        if recherche.strip():
            vue = vue[vue["PartNumber"].str.contains(recherche.strip(), case=False, regex=False)]
        if seulement:
            vue = vue[vue["Ecart"].abs() > 1e-9]
        st.caption(f"{len(vue):,} ligne(s) affichée(s) sur {len(df):,}".replace(",", " "))
        st.dataframe(
            vue,
            hide_index=True,
            column_config={
                "PU": st.column_config.NumberColumn("PU", format="%.2f"),
                "Ecart €": st.column_config.NumberColumn("Ecart €", format="%.2f"),
            },
        )

    with t_top:
        avec = df[df["Ecart"].abs() > 1e-9].copy()
        if avec.empty:
            st.success("✅ Aucun écart : le stock final correspond exactement aux mouvements.")
        else:
            avec["abs"] = avec["Ecart €"].abs()
            top = avec.sort_values("abs", ascending=False).head(15)
            st.caption("Les 15 articles avec le plus grand écart en valeur (€). Écart > 0 : excédent ; < 0 : manque.")
            st.bar_chart(top.set_index("PartNumber")["Ecart €"])
            st.dataframe(
                top[["PartNumber", "Type", "PU", "SI", "SF", "Ecart", "Ecart €"]],
                hide_index=True,
                column_config={
                    "PU": st.column_config.NumberColumn("PU", format="%.2f"),
                    "Ecart €": st.column_config.NumberColumn("Ecart €", format="%.2f"),
                },
            )

    with t_tcd:
        st.caption("Quantités signées par type de mouvement (colonnes), avec le total par PartNumber.")
        for i, (tcd, nom) in enumerate(zip(res["tcds"], res["noms_mouvements"]), start=1):
            with st.expander(f"TCD{i} — {nom}", expanded=(i == 1)):
                st.dataframe(tcd, hide_index=True)

    with t_ctrl:
        st.markdown("**Fichiers de mouvements**")
        st.dataframe(res["controle"], hide_index=True)

        st.markdown("**Magasins supprimés**")
        if res["magasins_supprimes"]:
            st.write(", ".join(res["magasins_supprimes"]))
        else:
            st.caption("Aucun magasin supprimé.")

        st.markdown("**Mouvements absents de la table des signes (signe 0 appliqué)**")
        if len(res["inconnus"]):
            st.warning(
                "Ces types de mouvement ne sont pas dans la table du plant : ils sont ignorés dans le calcul. "
                "S'ils comptent, ajoutez-les dans SIGNES_PAR_PLANT (app.py)."
            )
            st.dataframe(res["inconnus"], hide_index=True)
        else:
            st.success("Tous les mouvements sont reconnus.")

        st.markdown("**PartNumbers avec mouvements mais absents des deux stocks**")
        if len(res["orphelins"]):
            st.warning(
                "Ces articles ne figurent ni dans le stock initial ni dans le stock final : "
                "leurs mouvements ne sont pas dans la réconciliation."
            )
            st.dataframe(res["orphelins"], hide_index=True)
        else:
            st.success("Tous les articles avec mouvements existent dans les stocks.")

        if res["csv_redresses"]:
            st.markdown("**Fichiers de mouvements redressés (CW)**")
            for i, (nom, contenu) in enumerate(res["csv_redresses"]):
                base = nom.rsplit(".", 1)[0]
                st.download_button(
                    f"⬇️ {base}_CW_REDRESSE.csv",
                    data=contenu,
                    file_name=f"{base}_CW_REDRESSE.csv",
                    mime="text/csv",
                    key=f"dl_redresse_{i}",
                )


def main():
    st.set_page_config(page_title="Réconciliation de stock · LEONI", page_icon="📦", layout="wide")
    st.markdown(CSS, unsafe_allow_html=True)
    st.markdown(HERO, unsafe_allow_html=True)

    with st.sidebar:
        st.markdown("### Centre de contrôle")
        st.markdown(
            "1. Choisissez le **plant**\n"
            "2. Déposez le **stock initial** et le **stock final** (FMALS015)\n"
            "3. Déposez **1 ou 2 fichiers de mouvements** (FMALS032)\n"
            "4. Cochez les **magasins à supprimer** (si besoin)\n"
            "5. Lancez la réconciliation et téléchargez l'Excel"
        )
        st.divider()
        st.markdown("**Redressement CW** : détecté et appliqué automatiquement, fichier par fichier.")
        st.markdown("**Mouvement inconnu** : signe 0 (signalé dans l'onglet Contrôles).")
        st.divider()
        st.caption(
            "Les fichiers sont traités en mémoire : rien n'est écrit sur disque. "
            "Le cache des fichiers analysés expire après 15 minutes (2 entrées maximum par type). "
            "Les résultats restent en mémoire pour permettre l'affichage et le téléchargement."
        )

    st.markdown('<div class="workflow-kicker">01 — CONFIGURER & IMPORTER</div>', unsafe_allow_html=True)

    # ---------------------------------------------------------------- 1. PLANT
    with st.container(border=True):
        st.markdown("#### 1 · Plant")
        plant = st.selectbox(
            "Plant à réconcilier", PLANTS, index=None, placeholder="Choisir un plant…", key="plant"
        )
        if plant:
            with st.expander(f"Voir la table des signes du plant {plant}"):
                st.dataframe(table_signes(plant), hide_index=True)
        else:
            st.caption("Le plant détermine la table des signes appliquée aux mouvements.")

    # ---------------------------------------------------------------- 2. STOCKS
    with st.container(border=True):
        st.markdown("#### 2 · Stocks initial et final (FMALS015)")
        c1, c2 = st.columns(2)
        f_si = c1.file_uploader("Stock initial", type=["csv", "txt"], key="up_si")
        f_sf = c2.file_uploader("Stock final", type=["csv", "txt"], key="up_sf")

        if f_si is not None and f_sf is not None and f_si.getvalue() == f_sf.getvalue():
            st.warning("⚠️ Le stock initial et le stock final sont le même fichier.")
        if plant:
            for etiquette, f in (("initial", f_si), ("final", f_sf)):
                p = plant_dans_nom(f.name) if f is not None else None
                if p and p != plant:
                    st.warning(
                        f"⚠️ Le nom du stock {etiquette} ({f.name}) indique le plant {p}, "
                        f"mais le plant {plant} est sélectionné."
                    )

    # ---------------------------------------------------------------- 3. MOUVEMENTS
    mouvs, erreur_mouvements = [], False
    with st.container(border=True):
        st.markdown("#### 3 · Fichiers de mouvements (FMALS032)")
        f_mouv = st.file_uploader(
            "Un ou deux fichiers", type=["csv", "txt"], accept_multiple_files=True, key="up_mouv"
        )
        if f_mouv and len(f_mouv) > 2:
            st.error("❌ Maximum 2 fichiers de mouvements. Retirez-en pour continuer.")
            erreur_mouvements = True
        elif f_mouv:
            for f in sorted(f_mouv, key=lambda x: x.name):
                try:
                    m = parse_mouvements(f.getvalue(), f.name)
                except Exception as e:  # message lisible pour l'utilisateur
                    st.error(f"❌ {f.name} : {e}")
                    erreur_mouvements = True
                    continue
                mouvs.append(m)
                if m["redresse"]:
                    st.success(
                        f"✅ {m['nom']} — {_fmt(m['nb_lues'], 0)} lignes · "
                        f"CW détecté (ligne(s) concernée(s) : {_fmt(m['nb_cw'], 0)}) → "
                        "redressement appliqué automatiquement"
                    )
                else:
                    st.success(
                        f"✅ {m['nom']} — {_fmt(m['nb_lues'], 0)} lignes · "
                        "aucun CW détecté → redressement non nécessaire"
                    )

    # ---------------------------------------------------------------- 4. MAGASINS
    codes: set[str] = set()
    libelles: dict = {}
    with st.container(border=True):
        st.markdown("#### 4 · Magasins à supprimer")
        if not mouvs or erreur_mouvements:
            st.caption("La liste des magasins apparaît dès que les fichiers de mouvements sont chargés.")
        else:
            try:
                tableau = lister_magasins([m["df"] for m in mouvs])
            except Exception as e:
                st.error(f"❌ Liste des magasins impossible : {e}")
                tableau = None
                erreur_mouvements = True

            if tableau is not None:
                if tableau.empty:
                    st.info("Aucun magasin trouvé dans les fichiers de mouvements.")
                else:
                    st.caption(
                        f"{len(tableau)} magasin(s) trouvé(s) dans les mouvements (après redressement CW). "
                        "Cochez ceux à supprimer ; ne cochez rien pour tout conserver."
                    )
                    base = tableau.copy()
                    base.insert(0, "Supprimer", False)
                    cle = "mag_" + "|".join(str(_fid(f)) for f in sorted(f_mouv, key=lambda x: x.name))
                    edite = st.data_editor(
                        base,
                        key=cle,
                        hide_index=True,
                        disabled=["Magasin", "Lignes"],
                        column_config={
                            "Supprimer": st.column_config.CheckboxColumn("Supprimer", default=False),
                            "Magasin": st.column_config.TextColumn("Magasin"),
                            "Lignes": st.column_config.NumberColumn("Lignes", format="%d"),
                        },
                    )
                    cochees = edite.index[edite["Supprimer"].to_numpy()]
                    codes = set(cochees)
                    libelles = tableau["Magasin"].to_dict()
                    if codes:
                        nb_lignes = int(tableau.loc[list(codes), "Lignes"].sum())
                        st.info(
                            f"🗑️ {len(codes)} magasin(s) coché(s) — {_fmt(nb_lignes, 0)} ligne(s) "
                            "seront supprimées des mouvements."
                        )
                    else:
                        st.caption("Aucun magasin coché : aucune ligne ne sera supprimée.")

    st.markdown('<div class="workflow-kicker">02 — VALIDER & EXÉCUTER</div>', unsafe_allow_html=True)

    # ---------------------------------------------------------------- LANCEMENT
    manques = []
    if not plant:
        manques.append("choisir le plant")
    if f_si is None:
        manques.append("stock initial")
    if f_sf is None:
        manques.append("stock final")
    if not f_mouv:
        manques.append("fichier(s) de mouvements")
    if manques:
        st.info("À compléter : " + " · ".join(manques))

    pret = not manques and not erreur_mouvements and bool(mouvs)

    signature = None
    if pret:
        signature = (
            plant, _fid(f_si), _fid(f_sf),
            tuple(_fid(f) for f in sorted(f_mouv, key=lambda x: x.name)),
            tuple(sorted(codes)),
        )

    lancer = st.button("LANCER LA RÉCONCILIATION  →", type="primary", disabled=not pret, use_container_width=True)

    if lancer and pret:
        try:
            with st.spinner("Calcul en cours…"):
                df_si = parse_stock(f_si.getvalue(), f_si.name)
                df_sf = parse_stock(f_sf.getvalue(), f_sf.name)
                res = calculer(plant, df_si, f_si.name, df_sf, f_sf.name, mouvs, codes, libelles)
                # Les DataFrames de stock ne sont plus nécessaires localement :
                # les résultats requis restent dans res pour l'affichage et l'export.
                del df_si, df_sf
                gc.collect()
            res["signature"] = signature
            st.session_state["res"] = res
        except Exception as e:
            st.session_state.pop("res", None)
            st.error(f"❌ Calcul impossible : {e}")

    res = st.session_state.get("res")
    if res is not None:
        st.divider()
        if signature is not None and res.get("signature") == signature:
            afficher_resultats(res)
        else:
            st.info(
                "Les fichiers ou les paramètres ont changé depuis le dernier calcul : "
                "relancez la réconciliation pour mettre les résultats à jour."
            )


if __name__ == "__main__":
    main()
